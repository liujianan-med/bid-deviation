# -*- coding: utf-8 -*-
"""
01_read_params.py —— 读取招标参数表，抽出「编号 / 标记 / 需求描述」

这是你的第一个真实脚本。它不用大模型，只做一件事：
把 Excel 里 36 行参数读出来，结构化，存成 JSON。

运行方式（在项目根目录）：
    .venv\Scripts\python.exe scripts\01_read_params.py

你要理解的三件事：
    1. 程序 = 读数据 -> 处理 -> 写数据
    2. openpyxl 把 Excel 当成"格子矩阵"，cell(行, 列)，行列都从 1 开始
    3. 真实数据永远是脏的，脚本要能容忍脏
"""

import json
import re
from pathlib import Path

import openpyxl

# ---------- 路径：用 __file__ 定位，不要写死绝对路径 ----------
# __file__ 是当前脚本的路径；.parent 是 scripts/，再 .parent 是项目根目录
BASE_DIR = Path(__file__).resolve().parent.parent
SRC_FILE = BASE_DIR / "data" / "招标参数模版-皮秒治疗仪.xlsx"
OUT_FILE = BASE_DIR / "data" / "params.json"

# 这份表的结构：第 3 行是表头，第 4 行是空行（合并单元格留下的），第 5 行开始是数据
HEADER_ROW = 3
FIRST_DATA_ROW = 5


def clean(value):
    """把单元格的值变成干净的字符串。None 会变成空串，换行会变成空格。"""
    if value is None:
        return ""
    return str(value).replace("\n", " ").replace("\r", " ").strip()


def split_mark(text):
    """
    从 '★3.6 峰值功率...' 这种内容里拆出星级标记。

    ★ = 关键参数，不满足直接废标
    ▲ = 重要参数，影响评分
    这两个符号在标书里是有法律含义的，必须单独提出来，不能当普通字符丢掉。
    """
    mark = ""
    text = text.strip()
    for symbol in ("★", "▲"):
        if text.startswith(symbol):
            mark = symbol
            text = text[len(symbol):].strip()
            break
    return mark, text


def has_number(text):
    """判断这段要求里有没有具体数值（有数值才可能做量化比对）。"""
    return bool(re.search(r"\d", text))


def main():
    if not SRC_FILE.exists():
        print(f"[错误] 找不到文件：{SRC_FILE}")
        return

    print(f"读取：{SRC_FILE.name}")

    # data_only=True 表示读公式计算后的值，而不是公式本身
    wb = openpyxl.load_workbook(SRC_FILE, data_only=True)
    ws = wb["Sheet1"]

    # 先看看表头长什么样，确认列的顺序没变
    headers = [clean(ws.cell(HEADER_ROW, c).value) for c in range(1, 6)]
    print("表头：", headers)
    print("-" * 68)

    records = []       # 存放所有参数条目
    sections = []      # 存放章节标题行，比如「3.技术参数和性能参数」
    seen_ids = {}      # 用来发现重复编号

    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        raw_id = clean(ws.cell(row, 1).value)      # A 列：序号
        raw_req = clean(ws.cell(row, 2).value)     # B 列：科室需求（也是伟思对标列）

        # 整行都空 -> 跳过（表里确实有空行）
        if not raw_id and not raw_req:
            continue

        mark, item_id = split_mark(raw_id)

        # 这种情况是「章节标题」：序号有值，但需求描述是空的
        # 例如『3.技术参数和性能参数』下面挂着一堆 3.1 / 3.2
        if raw_req == "":
            sections.append(item_id)
            continue

        # 标记有时写在需求描述开头，有时写在序号列，两处都查一下
        if not mark:
            mark, raw_req = split_mark(raw_req)

        record = {
            "row": row,                 # 原 Excel 行号，方便回溯
            "item_id": item_id,         # 如 "3.6"
            "mark": mark,               # "★" / "▲" / ""
            "requirement": raw_req,     # 需求原文
            "quantifiable": has_number(raw_req),   # 是否含数值
            # 竞品列先原样留着，后面做对标时才用
            "competitor_ky": clean(ws.cell(row, 3).value),    # 科英
            "competitor_bd": clean(ws.cell(row, 4).value),    # 半岛
            "competitor_ad": clean(ws.cell(row, 5).value),    # 安德
        }
        records.append(record)

        if item_id in seen_ids:
            record["duplicate_id"] = True
        else:
            seen_ids[item_id] = row

    # ---------- 统计 ----------
    star = sum(1 for r in records if r["mark"] == "★")
    tri = sum(1 for r in records if r["mark"] == "▲")
    quant = sum(1 for r in records if r["quantifiable"])
    dup = [r for r in records if r.get("duplicate_id")]

    print(f"参数条目      : {len(records)} 条")
    print(f"章节标题行    : {len(sections)} 条  -> {sections}")
    print(f"★ 关键参数    : {star} 条（不满足即废标）")
    print(f"▲ 重要参数    : {tri} 条")
    print(f"含数值可量化  : {quant} 条（占 {quant / len(records) * 100:.0f}%）")
    print(f"编号重复      : {len(dup)} 条")

    if dup:
        print("\n[警告] 以下编号在原表里重复出现，做比对时会错位，需要人工核对：")
        for r in dup:
            print(f"  第 {r['row']} 行  编号 {r['item_id']}  <- 与前面重复")
            print(f"      {r['requirement'][:50]}")

    print("\n" + "-" * 68)
    print("前 5 条预览：")
    for r in records[:5]:
        flag = r["mark"] or " "
        print(f"  [{flag}] {r['item_id']:<6} {r['requirement'][:52]}")

    # ---------- 写结果 ----------
    payload = {
        "source": SRC_FILE.name,
        "headers": headers,
        "sections": sections,
        "count": len(records),
        "stats": {
            "star": star,
            "triangle": tri,
            "quantifiable": quant,
            "duplicate_id": len(dup),
        },
        "items": records,
    }

    OUT_FILE.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"\n已保存 -> {OUT_FILE}")
    print(f"文件大小 {OUT_FILE.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    # 这一行是 Python 的约定：只有直接运行这个文件时才执行 main()
    # 如果被别的文件 import，就不会自动跑
    main()
