# -*- coding: utf-8 -*-
"""
05_parse_photon.py —— 解析光子招标参数表，做成困难样本集

为什么换一张表？
    皮秒那张表上模型 100% 全对，但样本都来自同一份文件、同一种写法。
    光子表的表达方式不同（★ 标在文字里、更多功能描述、出现温度/时间等
    新单位），这才是对「泛化能力」的真正测试。

    注意这张表的结构和皮秒表不一样：
      皮秒表：序号列是 "3.1" 这种编号，★/▲ 是单独一列
      光子表：序号列是纯数字，★ 直接写在需求文字前面
    所以解析代码要单独写——「每份招标文件格式都不一样」本身就是
    这个项目要解决的现实问题。

运行：
    .venv\\Scripts\\python.exe scripts\\05_parse_photon.py
"""

import json
import os

import openpyxl

SRC = r"C:/Users/Administrator/Desktop/Excel/招标参数模版-光子治疗仪.xlsx"
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "params_photon.json")


def main():
    ws = openpyxl.load_workbook(SRC, data_only=True)["Sheet1"]

    items = []
    star = 0
    for r in range(5, 40):
        item_id = ws.cell(r, 1).value
        text = ws.cell(r, 2).value
        if not text:
            continue
        text = str(text).strip()

        # ★ 藏在文字里，要单独剥出来
        mark = None
        if text.startswith("★"):
            mark = "★"
            text = text.lstrip("★").strip()
            star += 1

        items.append({"row": r, "item_id": str(item_id).strip(), "mark": mark,
                      "requirement": text})

    data = {
        "source": "招标参数模版-光子治疗仪.xlsx",
        "headers": ["序号", "科室提报技术需求", "科英", "奇致", "金达威"],
        "count": len(items),
        "stats": {"star": star},
        "items": items,
    }
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"解析完成：{len(items)} 条，其中 ★ {star} 条 -> data/params_photon.json")
    for it in items:
        m = it["mark"] or " "
        print(f"  行{it['row']:>2} {m} [{it['item_id']:>4}] {it['requirement'][:50]}")


if __name__ == "__main__":
    main()
